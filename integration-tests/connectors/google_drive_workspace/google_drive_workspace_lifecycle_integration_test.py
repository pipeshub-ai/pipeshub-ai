# pyright: ignore-file

"""Google Drive Workspace — item lifecycle against a real tenant.

Covers the GD_Testing cases the folder-filter and entity suites do not: trash and
restore (with the vector store checked, not just the graph), permanent folder
delete, folder rename, same-named siblings, sharing with someone outside the
Workspace (and that person signing up), and a Google Doc too large for files.export.

One connector, ``folder_ids=[lifecycle root]`` with indexing on, shared by the
whole module (``drive_workspace_lifecycle_connector``). Tests run in order; each
leaves the tree in a state the next one expects.
"""

from __future__ import annotations

import io
import logging
import os
import sys
import uuid
from pathlib import Path
from typing import Any

import pytest

_ROOT = Path(__file__).resolve().parents[2]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from googleapiclient.errors import HttpError  # type: ignore[import-not-found]  # noqa: E402

from app.config.constants.arangodb import MimeTypes  # type: ignore[import-not-found]  # noqa: E402
from app.connectors.sources.google.drive.utils.drive_mime import (  # type: ignore[import-not-found]  # noqa: E402
    DOCX_MIME,
)
from app.sources.external.google.drive.drive import (  # type: ignore[import-not-found]  # noqa: E402
    GoogleDriveDataSource,
)
from connectors.google_drive_workspace.drive_workspace_test_utils import (  # noqa: E402
    delete_drive_item_permanently,
    ensure_pipeshub_user_exists,
    ensure_pipeshub_user_inactive,
    record_absent,
    record_present,
    rename_drive_item,
    share_drive_item_with_user,
    sync_until,
    trash_drive_item,
    upload_drive_file_resumable,
)
from helper.clients.users_client import UsersClient  # noqa: E402
from helper.graph_provider import GraphProviderProtocol  # noqa: E402
from helper.graph_provider_utils import async_poll_until  # noqa: E402
from helper.indexing_wait import wait_until_record_indexing_completed  # noqa: E402
from helper.second_user import log_in_existing_user, log_out_existing_user  # noqa: E402
from helper.vector_store import VectorStoreProbe  # noqa: E402
from pipeshub_client import PipeshubClient  # noqa: E402

logger = logging.getLogger("drive-workspace-lifecycle-test")

pytestmark = [
    pytest.mark.integration,
    pytest.mark.google_drive_workspace,
    pytest.mark.asyncio(loop_scope="session"),
]

ENV_OUTSIDE_USER = "GOOGLE_DRIVE_USER_OUTSIDE_WORKSPACE"
_INDEXING_WAIT_SEC = int(os.getenv("GOOGLE_DRIVE_WORKSPACE_INDEXING_WAIT", "600"))
_PERSON_WAIT_SEC = int(os.getenv("GOOGLE_DRIVE_WORKSPACE_USER_GRAPH_TIMEOUT", "120"))

# Google refuses files.export above 10 MB; four 1600 px noise images (~30 MB as
# docx) stay well over it after Drive converts the upload to a Google Doc.
_BIG_DOC_IMAGES = 4
_BIG_DOC_IMAGE_PX = 1600
_EXPORT_LIMIT_BYTES = 10 * 1024 * 1024


def _ids(state: dict[str, Any], *keys: str) -> list[str]:
    return [state[k] for k in keys]


def _all(graph_provider: GraphProviderProtocol, connector_id: str, external_ids: list[str], *, present: bool):
    checks = [
        (record_present if present else record_absent)(graph_provider, connector_id, eid)
        for eid in external_ids
    ]

    async def _check() -> bool:
        for check in checks:
            if not await check():
                return False
        return True

    return _check


class TestDriveWorkspaceLifecycle:
    """Trash, restore, delete, rename and sharing on one live connector."""

    @pytest.mark.order(1)
    async def test_tc_drv_lc_001_trash_removes_record_and_embeddings(
        self,
        drive_workspace_lifecycle_connector: dict[str, Any],
        drive_workspace_datasource: GoogleDriveDataSource,
        pipeshub_client: PipeshubClient,
        graph_provider: GraphProviderProtocol,
        vector_store: VectorStoreProbe,
    ) -> None:
        """TC-DRV-LC-001 (GD N19): trashing a file deletes its record and its vectors."""
        state = drive_workspace_lifecycle_connector
        connector_id = state["connector_id"]
        file_id = state["trash_file_id"]

        record = await wait_until_record_indexing_completed(
            graph_provider, connector_id, file_id,
            timeout=_INDEXING_WAIT_SEC, pipeshub_client=pipeshub_client,
            description="trash-me.txt indexed before trash",
        )
        assert record.virtual_record_id, "indexed record has no virtualRecordId"
        await vector_store.assert_embeddings_present(record.virtual_record_id)

        await trash_drive_item(drive_workspace_datasource, file_id)
        await sync_until(
            pipeshub_client, graph_provider, connector_id,
            record_absent(graph_provider, connector_id, file_id),
            description=f"trash-me.txt ({file_id}) removed after trash",
        )
        await vector_store.assert_embeddings_gone(record.virtual_record_id)
        state["trash_virtual_record_id"] = record.virtual_record_id

    @pytest.mark.order(2)
    async def test_tc_drv_lc_002_restore_resyncs_record(
        self,
        drive_workspace_lifecycle_connector: dict[str, Any],
        drive_workspace_datasource: GoogleDriveDataSource,
        pipeshub_client: PipeshubClient,
        graph_provider: GraphProviderProtocol,
    ) -> None:
        """TC-DRV-LC-002 (GD N20): restoring from trash syncs (and indexes) the file again."""
        state = drive_workspace_lifecycle_connector
        connector_id = state["connector_id"]
        file_id = state["trash_file_id"]
        if await graph_provider.get_record_by_external_id(connector_id, file_id) is not None:
            pytest.skip("TC-DRV-LC-001 did not leave the file trashed")

        await trash_drive_item(drive_workspace_datasource, file_id, trashed=False)
        await sync_until(
            pipeshub_client, graph_provider, connector_id,
            record_present(graph_provider, connector_id, file_id),
            description=f"trash-me.txt ({file_id}) back after restore",
        )
        parent = await graph_provider.get_record_parent_external_id(connector_id, file_id)
        assert parent == state["root_folder_id"], (
            f"restored file parent {parent!r} != lifecycle root {state['root_folder_id']!r}"
        )
        await wait_until_record_indexing_completed(
            graph_provider, connector_id, file_id,
            timeout=_INDEXING_WAIT_SEC, pipeshub_client=pipeshub_client,
            description="trash-me.txt indexed again after restore",
        )

    @pytest.mark.order(3)
    async def test_tc_drv_lc_003_permanent_folder_delete_removes_descendants(
        self,
        drive_workspace_lifecycle_connector: dict[str, Any],
        drive_workspace_datasource: GoogleDriveDataSource,
        pipeshub_client: PipeshubClient,
        graph_provider: GraphProviderProtocol,
    ) -> None:
        """TC-DRV-LC-003 (GD N21): deleting a folder for good removes every descendant."""
        state = drive_workspace_lifecycle_connector
        connector_id = state["connector_id"]
        subtree = _ids(state, "doomed_folder_id", "doomed_sub_id", "doomed_a_id", "doomed_b_id")

        await delete_drive_item_permanently(drive_workspace_datasource, state["doomed_folder_id"])
        try:
            await sync_until(
                pipeshub_client, graph_provider, connector_id,
                _all(graph_provider, connector_id, subtree, present=False),
                description="doomed/ and all its descendants removed after permanent delete",
            )
        except TimeoutError:
            left = [
                eid for eid in subtree
                if await graph_provider.get_record_by_external_id(connector_id, eid) is not None
            ]
            raise AssertionError(
                f"after deleting doomed/ for good these records are still in the graph: {left}"
            ) from None

    @pytest.mark.order(4)
    async def test_tc_drv_lc_004_folder_rename_keeps_children(
        self,
        drive_workspace_lifecycle_connector: dict[str, Any],
        drive_workspace_datasource: GoogleDriveDataSource,
        pipeshub_client: PipeshubClient,
        graph_provider: GraphProviderProtocol,
    ) -> None:
        """TC-DRV-LC-004 (GD N22): renaming a folder keeps its children's ids and edges."""
        state = drive_workspace_lifecycle_connector
        connector_id = state["connector_id"]
        folder_id = state["renamable_folder_id"]
        child_id = state["renamable_child_id"]

        folder_before = await graph_provider.get_record_by_external_id(connector_id, folder_id)
        child_before = await graph_provider.get_record_by_external_id(connector_id, child_id)
        assert folder_before is not None and child_before is not None

        new_name = f"renamed-{uuid.uuid4().hex[:6]}"
        await rename_drive_item(drive_workspace_datasource, folder_id, new_name)

        async def _renamed() -> bool:
            folder = await graph_provider.get_record_by_external_id(connector_id, folder_id)
            return folder is not None and folder.record_name == new_name

        await sync_until(
            pipeshub_client, graph_provider, connector_id, _renamed,
            description=f"folder {folder_id} renamed to {new_name}",
        )
        folder_after = await graph_provider.get_record_by_external_id(connector_id, folder_id)
        child_after = await graph_provider.get_record_by_external_id(connector_id, child_id)
        assert folder_after is not None and child_after is not None
        assert folder_after.id == folder_before.id, "folder record was replaced instead of renamed"
        assert child_after.id == child_before.id, "child record was replaced by the folder rename"
        child_parent = await graph_provider.get_record_parent_external_id(connector_id, child_id)
        assert child_parent == folder_id, f"child parent {child_parent!r} != renamed folder {folder_id!r}"

    @pytest.mark.order(5)
    async def test_tc_drv_lc_005_same_name_siblings_both_synced(
        self,
        drive_workspace_lifecycle_connector: dict[str, Any],
        graph_provider: GraphProviderProtocol,
    ) -> None:
        """TC-DRV-LC-005 (GD E2): two files with the same name in one folder are two records."""
        state = drive_workspace_lifecycle_connector
        connector_id = state["connector_id"]
        first = await graph_provider.get_record_by_external_id(connector_id, state["dup_first_id"])
        second = await graph_provider.get_record_by_external_id(connector_id, state["dup_second_id"])
        assert first is not None and second is not None, "a same-named sibling is missing"
        assert first.id != second.id, "same-named siblings collapsed into one record"
        assert first.record_name == second.record_name == "same-name.txt"
        for record_id in (state["dup_first_id"], state["dup_second_id"]):
            parent = await graph_provider.get_record_parent_external_id(connector_id, record_id)
            assert parent == state["dup_folder_id"], f"{record_id} parent {parent!r} != dup/"

    @pytest.mark.order(6)
    async def test_tc_drv_lc_006_outside_share_creates_person(
        self,
        drive_workspace_lifecycle_connector: dict[str, Any],
        drive_workspace_datasource: GoogleDriveDataSource,
        pipeshub_client: PipeshubClient,
        graph_provider: GraphProviderProtocol,
    ) -> None:
        """TC-DRV-LC-006 (GD N18): a share to an outside address becomes a Person with app access."""
        outside = os.getenv(ENV_OUTSIDE_USER, "").strip()
        if not outside:
            pytest.skip(f"{ENV_OUTSIDE_USER} not set")
        existing = await graph_provider.graph_find_user_by_email(outside)
        if existing is not None:
            # Signup (TC-DRV-LC-007) turns the Person into a User for good, so this case needs an
            # address that has never had an account in this org.
            pytest.skip(
                f"{outside} already has a PipesHub user in this org (signed up in an earlier run); "
                f"set {ENV_OUTSIDE_USER} to an outside address that never signed up"
            )
        state = drive_workspace_lifecycle_connector
        connector_id = state["connector_id"]
        file_id = state["share_out_file_id"]

        share_drive_item_with_user(drive_workspace_datasource, file_id, outside)

        async def _person_has_access() -> bool:
            label = await graph_provider.principal_record_permission_label(connector_id, outside, file_id)
            return label == "Person"

        await sync_until(
            pipeshub_client, graph_provider, connector_id, _person_has_access,
            description=f"Person {outside} holds PERMISSION on share-out.txt",
        )
        relation = await graph_provider.principal_app_relation(connector_id, outside)
        assert relation == {"label": "Person", "isExternalUser": True}, (
            f"outside person's app relation is {relation!r}"
        )
        state["outside_email"] = outside

    @pytest.mark.order(7)
    async def test_tc_drv_lc_007_outside_person_signup_becomes_user(
        self,
        drive_workspace_lifecycle_connector: dict[str, Any],
        users_client: UsersClient,
        graph_provider: GraphProviderProtocol,
    ) -> None:
        """TC-DRV-LC-007 (GD N18): when that outside person signs up, the Person becomes their User."""
        state = drive_workspace_lifecycle_connector
        outside = state.get("outside_email")
        if not outside:
            pytest.skip("TC-DRV-LC-006 did not create the outside Person")
        connector_id = state["connector_id"]
        file_id = state["share_out_file_id"]

        user_id, created = ensure_pipeshub_user_exists(users_client, outside)
        try:
            async def _converted() -> bool | None:
                label = await graph_provider.principal_record_permission_label(
                    connector_id, outside, file_id
                )
                return True if label == "User" else None

            await async_poll_until(
                _converted, timeout=_PERSON_WAIT_SEC, interval=3,
                description=f"{outside}'s PERMISSION moved from Person to User after signup",
            )
            relation = await graph_provider.principal_app_relation(connector_id, outside)
            assert relation is not None and relation["label"] == "User", (
                f"app relation after signup is {relation!r}"
            )
        finally:
            if created:
                ensure_pipeshub_user_inactive(users_client, outside)
                logger.info("Soft-deleted outside user %s (%s) created by this test", outside, user_id)

    @pytest.mark.order(8)
    async def test_tc_drv_lc_008_google_doc_over_export_limit_streams(
        self,
        drive_workspace_lifecycle_connector: dict[str, Any],
        drive_workspace_datasource: GoogleDriveDataSource,
        pipeshub_client: PipeshubClient,
        graph_provider: GraphProviderProtocol,
        tmp_path: Path,
    ) -> None:
        """TC-DRV-LC-008 (GD N16): a Google Doc over the 10 MB files.export limit still streams in full."""
        state = drive_workspace_lifecycle_connector
        connector_id = state["connector_id"]

        local = tmp_path / "big-export.docx"
        _write_noise_docx(local, images=_BIG_DOC_IMAGES, px=_BIG_DOC_IMAGE_PX)
        doc_id = await upload_drive_file_resumable(
            drive_workspace_datasource, local, state["root_folder_id"],
            name="big-export", mime_type=DOCX_MIME, convert_to_mime=MimeTypes.GOOGLE_DOCS.value,
        )
        # Precondition: without this the test would pass on a doc that never hit the limit.
        try:
            drive_workspace_datasource.client.files().export(fileId=doc_id, mimeType=DOCX_MIME).execute()
            pytest.fail("fixture doc exported under 10 MB; it no longer exercises the export limit")
        except HttpError as e:
            assert "exportSizeLimitExceeded" in str(e), f"unexpected export error: {e}"

        await sync_until(
            pipeshub_client, graph_provider, connector_id,
            record_present(graph_provider, connector_id, doc_id),
            description=f"big Google Doc ({doc_id}) synced",
        )
        record = await graph_provider.get_record_by_external_id(connector_id, doc_id)
        owner = log_in_existing_user(pipeshub_client, state["pipeshub_user_id"], state["test_user_email"])
        try:
            resp = owner.get(pipeshub_client._stream_record_path(record.id))
        finally:
            log_out_existing_user(pipeshub_client, owner)
        assert resp.status_code == 200, f"stream of over-limit Google Doc: HTTP {resp.status_code} {resp.text[:300]}"
        assert len(resp.content) > _EXPORT_LIMIT_BYTES, (
            f"streamed {len(resp.content)} bytes; expected the full export (> 10 MB)"
        )
        state["big_doc_id"] = doc_id

    @pytest.mark.order(8)
    @pytest.mark.xfail(
        strict=True,
        reason="B7: indexing stores the processed record through /api/v1/document/internal/upload, which "
        "reuses the org's user upload cap (default 30 MB); this doc's processed form is larger, so the "
        "upload is refused (400), retried as transient three times, and the record ends FAILED.",
    )
    async def test_tc_drv_lc_008b_google_doc_over_export_limit_indexes(
        self,
        drive_workspace_lifecycle_connector: dict[str, Any],
        pipeshub_client: PipeshubClient,
        graph_provider: GraphProviderProtocol,
    ) -> None:
        """TC-DRV-LC-008b (GD N16): that over-limit Google Doc also reaches indexingStatus COMPLETED."""
        state = drive_workspace_lifecycle_connector
        doc_id = state.get("big_doc_id")
        if not doc_id:
            pytest.skip("TC-DRV-LC-008 did not sync the over-limit Google Doc")
        await wait_until_record_indexing_completed(
            graph_provider, state["connector_id"], doc_id,
            timeout=_INDEXING_WAIT_SEC, pipeshub_client=pipeshub_client,
            description="over-limit Google Doc indexed",
        )

    @pytest.mark.order(8)
    async def test_tc_drv_lc_008c_refused_record_fails_once_with_the_reason(
        self,
        drive_workspace_lifecycle_connector: dict[str, Any],
        graph_provider: GraphProviderProtocol,
    ) -> None:
        """TC-DRV-LC-008c: a processed record the storage service refuses fails with that reason, not after retries."""
        state = drive_workspace_lifecycle_connector
        doc_id = state.get("big_doc_id")
        if not doc_id:
            pytest.skip("TC-DRV-LC-008 did not sync the over-limit Google Doc")
        connector_id = state["connector_id"]

        async def _settled() -> Any:
            record = await graph_provider.get_record_by_external_id(connector_id, doc_id)
            if record is None or record.indexing_status in (None, "QUEUED", "IN_PROGRESS", "NOT_STARTED"):
                return None
            return record

        record = await async_poll_until(
            _settled, timeout=_INDEXING_WAIT_SEC, interval=10,
            description="over-limit Google Doc reaches a final indexing status",
        )
        if record.indexing_status == "COMPLETED":
            pytest.skip("indexed: the storage cap no longer refuses this document")
        assert record.indexing_status == "FAILED", f"unexpected status {record.indexing_status!r}"
        reason = record.reason or ""
        assert "too large" in reason, f"the FAILED reason should say the file is too large, got {reason!r}"


def _write_noise_docx(path: Path, *, images: int, px: int) -> None:
    """A docx of incompressible images, so its Google Doc export cannot shrink under 10 MB."""
    import docx  # type: ignore[import-not-found]
    import numpy as np
    from docx.shared import Inches  # type: ignore[import-not-found]
    from PIL import Image

    document = docx.Document()
    document.add_paragraph(
        "Over-limit export fixture. The lighthouse keeper logs lamp hours every night."
    )
    rng = np.random.default_rng(7)
    for index in range(images):
        buffer = io.BytesIO()
        Image.fromarray(rng.integers(0, 255, (px, px, 3), dtype=np.uint8)).save(buffer, "PNG")
        buffer.seek(0)
        document.add_picture(buffer, width=Inches(6))
        document.add_paragraph(f"Figure {index}: lamp hours")
    document.save(path)
