"""Excel parser: gradient-filled cells in a hand-written (Excel-style) workbook.

openpyxl-built workbooks only exercise the fills openpyxl itself writes. This
file builds a small package whose styles.xml is written the way Excel writes
it (theme stops with a tint, path gradients, an indexed first stop, a gradient
with no stops) and sends it through the real entry points.
"""

import io
import json
import zipfile
from unittest.mock import MagicMock

import pytest
from langchain_core.messages import AIMessage
from openpyxl import Workbook, load_workbook

import app.modules.parsers.excel.excel_parser as excel_module
from app.modules.parsers.excel.excel_parser import ExcelParser
from app.modules.parsers.excel.prompt_template import ExcelHeaderDetection

_STYLES_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<fonts count="2">
<font><sz val="11"/><color theme="1"/><name val="Calibri"/></font>
<font><b/><sz val="11"/><color rgb="FFC00000"/><name val="Calibri"/></font>
</fonts>
<fills count="6">
<fill><patternFill patternType="none"/></fill>
<fill><patternFill patternType="gray125"/></fill>
<fill><gradientFill degree="90"><stop position="0"><color theme="4" tint="-0.25"/></stop><stop position="1"><color theme="0"/></stop></gradientFill></fill>
<fill><gradientFill type="path" left="0.5" right="0.5" top="0.5" bottom="0.5"><stop position="0"><color rgb="FFFF0000"/></stop><stop position="1"><color rgb="FF0000FF"/></stop></gradientFill></fill>
<fill><gradientFill><stop position="0"><color indexed="64"/></stop><stop position="1"><color rgb="FF00B050"/></stop></gradientFill></fill>
<fill><gradientFill degree="180"/></fill>
</fills>
<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
<cellXfs count="5">
<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>
<xf numFmtId="0" fontId="1" fillId="2" borderId="0" xfId="0" applyFont="1" applyFill="1"/>
<xf numFmtId="0" fontId="0" fillId="3" borderId="0" xfId="0" applyFill="1"/>
<xf numFmtId="0" fontId="0" fillId="4" borderId="0" xfId="0" applyFill="1"/>
<xf numFmtId="0" fontId="0" fillId="5" borderId="0" xfId="0" applyFill="1"/>
</cellXfs>
<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>
</styleSheet>"""


def _s(ref: str, style: int, text: str) -> str:
    return f'<c r="{ref}" s="{style}" t="inlineStr"><is><t>{text}</t></is></c>'


_SHEET_XML = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<dimension ref="A1:C4"/><sheetData>
<row r="1">{_s("A1", 1, "Name")}{_s("B1", 2, "Score")}{_s("C1", 0, "Group")}</row>
<row r="2">{_s("A2", 3, "alpha")}<c r="B2" s="4"><v>10</v></c>{_s("C2", 0, "g1")}</row>
<row r="3">{_s("A3", 2, "beta")}<c r="B3" s="2"/>{_s("C3", 0, "g2")}</row>
<row r="4">{_s("A4", 0, "gamma")}<c r="B4"><v>30</v></c>{_s("C4", 0, "g3")}</row>
</sheetData><mergeCells count="1"><mergeCell ref="A3:B3"/></mergeCells></worksheet>"""

_VALUES = ["Name", "Score", "alpha", "beta", "gamma", "g1", "g2", "g3"]


def _excel_style_xlsx() -> bytes:
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.active["A1"] = "x"
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.infolist():
            if item.filename == "xl/styles.xml":
                z.writestr(item, _STYLES_XML)
            elif item.filename == "xl/worksheets/sheet1.xml":
                z.writestr(item, _SHEET_XML)
            else:
                z.writestr(item, src.read(item.filename))
    return out.getvalue()


@pytest.fixture
def xlsx_bytes() -> bytes:
    data = _excel_style_xlsx()
    sheet = load_workbook(io.BytesIO(data))["Sheet1"]
    # The fixture is only meaningful if openpyxl really read gradient fills.
    assert {sheet[c].fill.tagname for c in ("A1", "B1", "A2", "B2", "A3")} == {"gradientFill"}
    return data


@pytest.fixture
def parser() -> ExcelParser:
    return ExcelParser(logger=MagicMock(), config_service=MagicMock())


@pytest.fixture
def stub_llm(monkeypatch: pytest.MonkeyPatch) -> None:
    class _Gateway:
        async def invoke(self, *args: object, **kwargs: object) -> AIMessage:
            return AIMessage(content="stub summary")

    async def structured(llm: object, messages: object, model: object, *a: object, **k: object) -> object:
        if model is ExcelHeaderDetection:
            return ExcelHeaderDetection(
                has_headers=True, num_header_rows=1, confidence="high", reasoning="stub"
            )
        return None

    async def rows(*args: object, **kwargs: object) -> None:
        return None

    async def get_llm(*args: object, **kwargs: object) -> tuple[MagicMock, dict]:
        return MagicMock(name="llm"), {}

    monkeypatch.setattr(excel_module, "invoke_with_structured_output_and_reflection", structured)
    monkeypatch.setattr(excel_module, "invoke_with_row_descriptions_and_reflection", rows)
    monkeypatch.setattr(excel_module, "get_llm_gateway", lambda: _Gateway())
    monkeypatch.setattr(excel_module, "provider_key", lambda llm: "stub")
    monkeypatch.setattr(excel_module, "get_llm_invoke_config", lambda llm: {})
    monkeypatch.setattr(excel_module, "get_llm_for_role", get_llm)


def _colours(parser: ExcelParser, xlsx: bytes) -> dict[str, tuple[str | None, str | None]]:
    parser.load_workbook_from_binary(xlsx)
    sheet = parser.workbook["Sheet1"]
    out = {}
    for coord in ("A1", "B1", "A2", "B2", "A3"):
        cell = sheet[coord]
        style = parser._process_cell(cell, None, cell.row, cell.column)["style"]
        out[coord] = (style["fill"]["background_color"], style["font"]["color"])
    return out


def test_excel_style_gradient_colours(parser: ExcelParser, xlsx_bytes: bytes) -> None:
    got = _colours(parser, xlsx_bytes)
    # theme first stop (with tint): no RGB, so None; the header font is a real RGB.
    assert got["A1"] == (None, "FFC00000")
    # path gradient: first stop's RGB; the default font is a theme colour: None.
    assert got["B1"] == ("FFFF0000", None)
    # indexed first stop: None, not a crash.
    assert got["A2"] == (None, None)
    # no stops at all: None.
    assert got["B2"] == (None, None)
    # merged top-left cell with a path gradient.
    assert got["A3"] == ("FFFF0000", None)


@pytest.mark.asyncio
async def test_excel_style_gradient_lightweight_and_workbook_paths(
    parser: ExcelParser, xlsx_bytes: bytes
) -> None:
    parser.load_workbook_from_binary(xlsx_bytes)
    light = await parser.create_blocks_lightweight()
    text = json.dumps([b.data for b in light.blocks], default=str)
    assert all(v in text for v in ("alpha", "beta", "gamma", "g1", "g3"))

    full = await ExcelParser(logger=MagicMock(), config_service=MagicMock()).parse_workbook(xlsx_bytes)
    text = json.dumps([b.data for b in full.blocks], default=str)
    assert all(v in text for v in ("alpha", "beta", "gamma"))


@pytest.mark.asyncio
async def test_excel_style_gradient_find_tables_colours_serialisable(
    parser: ExcelParser, xlsx_bytes: bytes, stub_llm: None
) -> None:
    parser.load_workbook_from_binary(xlsx_bytes)
    tables = await parser.find_tables(parser.workbook["Sheet1"], MagicMock())
    assert tables
    colours = [
        value
        for table in tables
        for row in table["data"]
        for cell in row
        for value in (
            cell["style"]["fill"].get("background_color"),
            cell["style"]["font"].get("color"),
        )
    ]
    assert all(c is None or isinstance(c, str) for c in colours)
    json.dumps(tables)  # plain JSON, no descriptor objects


@pytest.mark.asyncio
async def test_excel_style_gradient_parse_with_stubbed_llm(
    parser: ExcelParser, xlsx_bytes: bytes, stub_llm: None
) -> None:
    result = await parser.parse(xlsx_bytes, "styled.xlsx")
    container = result.block_container
    text = json.dumps(
        [b.data for b in container.blocks] + [g.data for g in container.block_groups],
        default=str,
    )
    assert all(v in text for v in ("alpha", "beta", "gamma", "g1", "g3"))
