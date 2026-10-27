"""A workbook too large to load is refused before openpyxl builds its cell objects."""

import io
import logging
import zipfile
from unittest.mock import MagicMock

import openpyxl
import pytest
from openpyxl import Workbook

from app.modules.parsers.excel import excel_parser as module
from app.modules.parsers.excel.excel_parser import ExcelParser
from app.services.parsing.interface import ParseError, ParseErrorCode
from app.utils import user_errors


def _workbook(rows: int, cols: int = 5, write_only: bool = False) -> bytes:
    wb = Workbook(write_only=write_only)
    ws = wb.create_sheet("data") if write_only else wb.active
    for r in range(rows):
        ws.append([f"v{r}c{c}" for c in range(cols)])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _with_dimension(content: bytes, ref: str) -> bytes:
    """The same workbook with every sheet's <dimension> replaced (or dropped when ref is empty)."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info)
            if info.filename.startswith("xl/worksheets/"):
                import re

                data = re.sub(rb"<dimension[^>]*/>", f'<dimension ref="{ref}"/>'.encode() if ref else b"", data)
            dst.writestr(info.filename, data)
    return out.getvalue()


@pytest.fixture
def parser() -> ExcelParser:
    return ExcelParser(logging.getLogger("t"), MagicMock())


@pytest.fixture
def load_spy(monkeypatch: pytest.MonkeyPatch) -> MagicMock:
    spy = MagicMock(side_effect=module.load_workbook)
    monkeypatch.setattr(module, "load_workbook", spy)
    return spy


def test_declared_cells_over_the_limit_are_refused_before_loading(
    parser: ExcelParser, load_spy: MagicMock, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(module, "EXCEL_MAX_CELLS", 100, raising=False)
    with pytest.raises(ParseError) as raised:
        parser.load_workbook_from_binary(_workbook(rows=50))  # 250 cells
    assert raised.value.code == ParseErrorCode.INVALID_INPUT
    assert raised.value.message == user_errors.EXCEL_TOO_LARGE
    load_spy.assert_not_called()


def test_a_missing_dimension_cannot_let_a_huge_sheet_through(
    parser: ExcelParser, load_spy: MagicMock, monkeypatch: pytest.MonkeyPatch
) -> None:
    content = _with_dimension(_workbook(rows=200), "")
    monkeypatch.setattr(module, "EXCEL_MAX_SHEET_XML_BYTES", 2000, raising=False)
    with pytest.raises(ParseError):
        parser.load_workbook_from_binary(content)
    load_spy.assert_not_called()


def test_an_understated_dimension_cannot_let_a_huge_sheet_through(
    parser: ExcelParser, load_spy: MagicMock, monkeypatch: pytest.MonkeyPatch
) -> None:
    content = _with_dimension(_workbook(rows=200), "A1")
    monkeypatch.setattr(module, "EXCEL_MAX_SHEET_XML_BYTES", 2000, raising=False)
    with pytest.raises(ParseError):
        parser.load_workbook_from_binary(content)
    load_spy.assert_not_called()


def test_a_malformed_dimension_is_ignored_not_trusted(
    parser: ExcelParser, load_spy: MagicMock, monkeypatch: pytest.MonkeyPatch
) -> None:
    content = _with_dimension(_workbook(rows=200), "not-a-range")
    monkeypatch.setattr(module, "EXCEL_MAX_SHEET_XML_BYTES", 2000, raising=False)
    with pytest.raises(ParseError):
        parser.load_workbook_from_binary(content)
    load_spy.assert_not_called()


def test_a_workbook_within_the_limits_loads_and_parses_as_before(parser: ExcelParser, load_spy: MagicMock) -> None:
    content = _workbook(rows=40)
    expected = ExcelParser(logging.getLogger("t"), MagicMock())._process_sheet(
        openpyxl.load_workbook(io.BytesIO(content), data_only=True).active
    )
    parser.load_workbook_from_binary(content)
    load_spy.assert_called_once()
    assert parser._process_sheet(parser.workbook.active) == expected


def test_the_refusal_reason_is_stored_as_written() -> None:
    error = ParseError(ParseErrorCode.INVALID_INPUT, user_errors.EXCEL_TOO_LARGE)
    assert user_errors.to_user_reason(error) == user_errors.EXCEL_TOO_LARGE


def test_a_file_that_is_not_a_zip_is_left_for_the_loader_to_report(parser: ExcelParser) -> None:
    with pytest.raises(Exception) as raised:
        parser.load_workbook_from_binary(b"not a workbook")
    assert not isinstance(raised.value, ParseError)
