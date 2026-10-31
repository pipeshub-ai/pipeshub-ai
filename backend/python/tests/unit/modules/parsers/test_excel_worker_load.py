# ruff: noqa: ANN202, DTZ001
"""Excel workbooks load in a memory-limited worker; the parent parses a compact model.

Covers the properties of the design, each of which fails without its part:
the worker refuses a load past its allowance (and the same file loads without the
limit); the parent never builds openpyxl objects; the model gives the output the
openpyxl path gives; a worker that dies leaves nothing behind; every caller goes
through the worker off the event loop. Real worker processes throughout.
"""
from __future__ import annotations

import asyncio
import datetime as dt
import io
import json
import logging
import os
import sys
import threading
import time
import zipfile
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

import openpyxl
import pytest
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Color, Font, GradientFill, PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink

from app.models.blocks import BlocksContainer
from app.modules.parsers import parse_pool
from app.modules.parsers.excel import excel_parser as module
from app.modules.parsers.excel.excel_parser import ExcelParser
from app.modules.parsers.excel.workbook_model import SheetModel, WorkbookModel
from app.services.parsing.interface import ParseError
from app.utils import user_errors
from tests.support.excel_workbooks import empty_cells, load, workbook_zip

MB = 1024 * 1024


def _realistic_workbook() -> bytes:
    wb = openpyxl.Workbook()
    sales = wb.active
    sales.title = "Sales"
    sales.append(["Region", "Units", "Price", "Shipped", "Time", "Active"])
    for i in range(1, 301):
        sales.append([f"R{i % 7}", i, i * 1.25, dt.datetime(2024, 1, 1) + dt.timedelta(days=i), dt.time(i % 24, i % 60), i % 2 == 0])
    sales["A1"].font = Font(bold=True, italic=True, size=14, color="FFC00000")
    sales["A1"].fill = PatternFill("solid", start_color="FFFFFF00")
    sales["B1"].fill = GradientFill(stop=("FF00FF00", "FF0000FF"))
    sales["C1"].font = Font(color=Color(theme=1))
    sales["D2"].number_format = "dd/mm/yyyy"
    sales["E3"].alignment = Alignment(horizontal="center", vertical="top")
    sales["F2"] = "=SUM(B2:B10)"  # no cached value: reads as empty
    sales["H1"] = "Notes"  # a second table to the right, after a blank column
    for r in range(2, 6):
        sales.cell(r, 8, f"note {r}")
    for r in range(310, 316):  # and one below, after a gap
        sales.cell(r, 1, f"x{r}")
        sales.cell(r, 2, r)

    report = wb.create_sheet("Report")
    report["A1"] = "Quarterly report"
    report.merge_cells("A1:C1")
    report["A2"] = "Region"
    report["B2"] = "Q1"
    report["C2"] = "Q2"
    for r in range(3, 20):
        report.append([f"R{r}", r, r * 2])
    report.merge_cells("A5:A7")
    report.merge_cells("B10:C11")
    report["B10"] = dt.datetime(2024, 5, 5, 13, 30)
    report["A9"].hyperlink = Hyperlink(ref="A9", location="Sales!A1")
    report["B3"].comment = Comment("hello", "me")

    wb.create_sheet("Empty")
    sparse = wb.create_sheet("Sparse")
    sparse["D10"] = "far"
    sparse["E10"] = 1
    sparse["D11"] = "away"
    sparse["E11"] = 2
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _parser() -> ExcelParser:
    return ExcelParser(logging.getLogger("t"), MagicMock())


def _stub_llm(parser: ExcelParser) -> None:
    def detect(rows, _llm):
        first = rows[0] if rows else []
        if first and first[0] == "Quarterly report":
            return SimpleNamespace(has_headers=True, num_header_rows=2, confidence="high")
        if first and all(isinstance(v, str) for v in first if v is not None):
            return SimpleNamespace(has_headers=True, num_header_rows=1, confidence="high")
        return SimpleNamespace(has_headers=False, num_header_rows=0, confidence="high")

    async def generate(sample_rows, column_count, _llm):
        return [f"Gen{i}" for i in range(column_count)]

    parser.detect_excel_headers_with_llm = AsyncMock(side_effect=detect)
    parser.generate_excel_headers_with_llm = AsyncMock(side_effect=generate)


def _dump(value: object) -> str:
    return json.dumps(value, default=str, sort_keys=True)


def _shape(container) -> dict:
    ids = {"__all__": {"id"}}
    return container.model_dump(mode="json", exclude={"blocks": ids, "block_groups": ids})


# -- the worker refuses a load past its allowance --------------------------------------------


def _heavy_workbook() -> bytes:
    # A 520 000-cell merged range: one MergedCell per cell once loaded, 60 bytes of XML.
    return workbook_zip('<sheetData/><mergeCells><mergeCell ref="A1:Z20000"/></mergeCells>')


def test_the_worker_refuses_a_load_past_its_allowance(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(module, "EXCEL_LOAD_MAX_BYTES", 32 * MB)
    with pytest.raises(ParseError) as raised:
        load(_heavy_workbook())
    assert raised.value.message == user_errors.EXCEL_TOO_LARGE


def test_the_same_workbook_loads_when_the_allowance_is_not_enforced(monkeypatch: pytest.MonkeyPatch) -> None:
    # Proves the refusal above comes from the worker's limit, not from the workbook being rejected.
    monkeypatch.setattr(module, "EXCEL_LOAD_MAX_BYTES", 4096 * MB)
    parser = load(_heavy_workbook())
    assert parser.workbook["S"].cell(2, 2).coordinate == "B2"
    assert parser.workbook.worker_peak_kb > 60 * 1024  # the 520 000 merged cells really were built, in the worker


def test_the_allowance_is_a_growth_cap_relative_to_the_worker_it_runs_in(monkeypatch: pytest.MonkeyPatch) -> None:
    """A long-lived pool worker that already holds memory still has the full allowance, and the
    limit is put back, so the next job (any parser) is not capped."""
    monkeypatch.setenv(parse_pool.WORKERS_ENV, "1")
    parse_pool.set_resource_governor(SimpleNamespace(ceilings=SimpleNamespace(heavy=4), report_memory_incident=lambda *_: None))
    monkeypatch.setattr(module, "EXCEL_LOAD_MAX_BYTES", 32 * MB)
    try:
        parse_pool.run_blocking(bytearray, 150 * MB)  # a previous job leaves the worker larger
        small = workbook_zip(f"<sheetData>{empty_cells(50, 10)}</sheetData>")
        assert load(small).workbook["S"].max_row == 50  # loads on the grown worker
        with pytest.raises(ParseError):
            load(_heavy_workbook())  # the same worker still refuses
        assert len(parse_pool.run_blocking(bytearray, 200 * MB)) == 200 * MB  # limit restored for the next job
    finally:
        parse_pool.shutdown_parse_pool()
        parse_pool.set_resource_governor(None)


# -- the parent never builds openpyxl objects --------------------------------------------------


@pytest.mark.asyncio
async def test_the_parent_never_builds_openpyxl_objects(monkeypatch: pytest.MonkeyPatch) -> None:
    data = _realistic_workbook_bytes_cached()  # built with openpyxl, before the parent's openpyxl is disabled
    import openpyxl.cell.cell as cell_module
    import openpyxl.reader.excel as reader
    import openpyxl.workbook.workbook as workbook_module

    def forbidden(*_a, **_k):
        raise AssertionError("openpyxl object built in the parent")

    for owner, name in [
        (reader, "load_workbook"),
        (reader.ExcelReader, "__init__"),
        (workbook_module.Workbook, "__init__"),
        (cell_module.Cell, "__init__"),
        (cell_module.MergedCell, "__init__"),
    ]:
        monkeypatch.setattr(owner, name, forbidden)
    parser = _parser()
    _stub_llm(parser)
    parser.load_workbook_from_binary(data)
    assert isinstance(parser.workbook, WorkbookModel)
    for name in parser.workbook.sheetnames:
        sheet = parser.workbook[name]
        assert isinstance(sheet, SheetModel)
        await parser.find_tables(sheet, MagicMock())
        parser._process_sheet(sheet)
    container = await parser.create_blocks_lightweight()
    assert container.blocks
    # and the module no longer imports the loader at all
    assert not hasattr(module, "load_workbook")


_CACHED: list[bytes] = []


def _realistic_workbook_bytes_cached() -> bytes:
    if not _CACHED:
        _CACHED.append(_realistic_workbook())
    return _CACHED[0]


# -- the model gives the output the openpyxl path gives ----------------------------------------


def _openpyxl_parser(data: bytes) -> ExcelParser:
    parser = _parser()
    parser.workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    return parser


def test_every_cell_of_every_sheet_is_processed_identically() -> None:
    data = _realistic_workbook_bytes_cached()
    expected = _openpyxl_parser(data)
    actual = load(data)
    assert actual.workbook.sheetnames == expected.workbook.sheetnames
    for name in expected.workbook.sheetnames:
        want = expected._process_sheet(expected.workbook[name])
        got = actual._process_sheet(actual.workbook[name])
        assert _dump(got) == _dump(want), name
        # random access, including merged and never-stored positions
        for row in range(1, expected.workbook[name].max_row + 3):
            for col in range(1, expected.workbook[name].max_column + 3):
                a = actual._process_cell(actual.workbook[name].cell(row, col), "h", row, col)
                b = expected._process_cell(expected.workbook[name].cell(row, col), "h", row, col)
                assert _dump(a) == _dump(b), (name, row, col)


@pytest.mark.asyncio
async def test_table_finding_and_the_basic_blocks_are_identical() -> None:
    data = _realistic_workbook_bytes_cached()
    expected = _openpyxl_parser(data)
    actual = load(data)
    _stub_llm(expected)
    _stub_llm(actual)
    for name in expected.workbook.sheetnames:
        want = await expected.find_tables(expected.workbook[name], MagicMock())
        got = await actual.find_tables(actual.workbook[name], MagicMock())
        assert _dump(got) == _dump(want), name
    seen = lambda p: [c.args[0] for c in p.detect_excel_headers_with_llm.await_args_list]  # noqa: E731
    assert seen(actual) == seen(expected)
    assert _shape(actual._build_basic_block_container()) == _shape(expected._build_basic_block_container())


@pytest.mark.asyncio
async def test_the_whole_llm_path_makes_identical_blocks() -> None:
    data = _realistic_workbook_bytes_cached()
    out = []
    for parser in (_openpyxl_parser(data), load(data)):
        _stub_llm(parser)
        parser.get_table_summary = AsyncMock(return_value="summary")
        parser.get_rows_text = AsyncMock(side_effect=lambda rows, _summary: [f"row {r[0]['row']}: " + ", ".join(str(c['value']) for c in r) for r in rows])
        parser.llm = MagicMock()
        with patch.object(module, "get_llm_gateway", MagicMock()), patch.object(parser, "process_sheet_with_summaries", wraps=parser.process_sheet_with_summaries):
            out.append(_shape(await parser.get_blocks_from_workbook(MagicMock())))
    assert out[0] == out[1]


def test_the_model_is_much_smaller_than_the_workbook_it_stands_for() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 3001):
        ws.append([r, f"name {r}", r * 1.5, dt.datetime(2024, 1, 1), None, True])
    out = io.BytesIO()
    wb.save(out)
    import pickle

    model = load(out.getvalue()).workbook
    sheet = model["Sheet"]
    assert len(sheet.keys) == 3000 * 5  # the empty column is never stored
    assert len(pickle.dumps(model)) < 18_000 * 60  # about 60 bytes a cell, openpyxl holds ~430


# -- a worker that dies leaves nothing behind -----------------------------------------------------


def _alive(pid: int) -> bool:
    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    try:
        with open(f"/proc/{pid}/stat") as stat:
            return stat.read().rsplit(")", 1)[1].split()[0] != "Z"
    except OSError:
        return False


def test_a_worker_killed_while_loading_is_a_retryable_crash_and_leaves_no_process(monkeypatch: pytest.MonkeyPatch) -> None:
    real = parse_pool.run_blocking
    child: list[int] = []

    def die_after_forking(_fn, *_args, **kwargs):
        code = "import os, time; os.closerange(3, 4096); time.sleep(600)"
        child.append(real(os.posix_spawn, sys.executable, [sys.executable, "-c", code], {}, **kwargs))
        return real(os._exit, 137, **kwargs)  # the kernel's OOM kill looks the same

    monkeypatch.setattr(parse_pool, "run_blocking", die_after_forking)
    # Nothing says the kill was for memory (an operator, a deploy), so it is the pool's
    # worker-crash reason, not "too large".
    with pytest.raises(parse_pool.ParseWorkerError) as raised:
        load(workbook_zip("<sheetData/>"))
    assert raised.value.cause == "crashed"
    assert str(raised.value) == user_errors.PARSE_WORKER_OUT_OF_MEMORY
    deadline = time.monotonic() + 10
    while _alive(child[0]) and time.monotonic() < deadline:
        time.sleep(0.05)
    assert not _alive(child[0])  # the worker's whole process group went with it


def test_the_parent_keeps_serving_after_a_worker_died() -> None:
    with pytest.raises(parse_pool.ParseWorkerError):
        parse_pool.run_blocking(os._exit, 137)
    assert load(workbook_zip(f"<sheetData>{empty_cells(3, 3, lambda ref: f'<c r={chr(34)}{ref}{chr(34)}><v>1</v></c>')}</sheetData>")).workbook["S"].max_row == 3


def test_a_worker_that_overruns_its_time_is_the_existing_timeout_reason() -> None:
    with pytest.raises(parse_pool.ParseWorkerError) as raised:
        parse_pool.run_blocking(time.sleep, 30, timeout=0.5)
    assert raised.value.cause == "timed_out"
    assert str(raised.value) == user_errors.PROCESSING_TIMED_OUT


def test_a_worker_that_cannot_start_is_never_replaced_by_loading_in_this_process(monkeypatch: pytest.MonkeyPatch) -> None:
    def cannot_start(*_a, **_k):
        raise parse_pool._WorkerUnavailable("no fork")

    monkeypatch.setattr(parse_pool, "_Worker", cannot_start)
    with pytest.raises(parse_pool.ParseWorkerError) as raised:
        load(workbook_zip("<sheetData/>"))
    assert raised.value.cause == "unavailable"


# -- every caller goes through the worker, off the event loop --------------------------------------


@pytest.fixture
def worker_calls(monkeypatch: pytest.MonkeyPatch) -> list[int]:
    threads: list[int] = []
    real = parse_pool.run_blocking

    def spy(fn, *args, **kwargs):
        threads.append(threading.get_ident())
        return real(fn, *args, **kwargs)

    monkeypatch.setattr(parse_pool, "run_blocking", spy)
    return threads


def _small_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    wb.active.append(["a", "b"])
    wb.active.append([1, 2])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.mark.asyncio
async def test_parse_loads_in_the_worker_off_the_event_loop(worker_calls: list[int]) -> None:
    parser = _parser()
    _stub_llm(parser)
    with patch.object(module, "get_llm_for_role", AsyncMock(return_value=(MagicMock(), {}))), patch.object(
        ExcelParser, "create_blocks", AsyncMock(return_value=BlocksContainer(blocks=[], block_groups=[]))
    ):
        await parser.parse(_small_xlsx(), "book.xlsx")
    assert worker_calls
    assert threading.get_ident() not in worker_calls


@pytest.mark.asyncio
async def test_parse_workbook_loads_in_the_worker_off_the_event_loop(worker_calls: list[int]) -> None:
    container = await _parser().parse_workbook(_small_xlsx())
    assert container.blocks
    assert threading.get_ident() not in worker_calls


@pytest.mark.asyncio
async def test_chat_attachments_load_in_the_worker_off_the_event_loop(worker_calls: list[int]) -> None:
    from app.api.routes.chatbot import _build_excel_blocks

    container = await _build_excel_blocks(_small_xlsx(), "book.xlsx", MagicMock())
    assert container.blocks
    assert threading.get_ident() not in worker_calls


@pytest.mark.asyncio
async def test_parse_file_loads_in_the_worker_off_the_event_loop(worker_calls: list[int]) -> None:
    from tests.unit.agents.actions.test_parse_file import _make_parser

    parser = _make_parser()
    with patch("app.agents.actions.util.parse_file.get_llm_for_role", AsyncMock(return_value=(MagicMock(), None))), patch.object(
        ExcelParser, "create_blocks", AsyncMock(return_value=MagicMock())
    ):
        await parser.handle_xlsx(_small_xlsx(), "book.xlsx")
    assert worker_calls
    assert threading.get_ident() not in worker_calls


@pytest.mark.asyncio
async def test_the_indexing_processor_loads_in_the_worker_off_the_event_loop(worker_calls: list[int]) -> None:
    from tests.unit.events.test_processor_extended import (
        _collect_events,
        _make_processor,
        _mock_record_dict,
    )

    real = _parser()
    proc = _make_processor(parsers={"xlsx": real})
    proc.graph_provider.get_document = AsyncMock(return_value=_mock_record_dict(recordName="test.xlsx"))
    with patch("app.events.processor.get_llm_for_role", new_callable=AsyncMock) as llm, patch(
        "app.events.processor.IndexingPipeline"
    ) as pipeline, patch.object(ExcelParser, "create_blocks", AsyncMock(return_value=MagicMock())):
        llm.return_value = (MagicMock(), {})
        pipeline.return_value.apply = AsyncMock()
        await _collect_events(
            proc.process_excel_document("test.xlsx", "r1", "1", "src", "o1", _small_xlsx(), "vr1", event_type="newRecord")
        )
    assert worker_calls
    assert threading.get_ident() not in worker_calls


@pytest.mark.asyncio
async def test_loading_does_not_block_the_event_loop() -> None:
    ticks = 0

    async def ticker() -> None:
        nonlocal ticks
        while True:
            await asyncio.sleep(0.01)
            ticks += 1

    task = asyncio.create_task(ticker())
    try:
        await asyncio.to_thread(load, _heavy_workbook_small_limit())
    except ParseError:
        pass
    finally:
        task.cancel()
    assert ticks >= 10


def _heavy_workbook_small_limit() -> bytes:
    return workbook_zip(f"<sheetData>{empty_cells(500, 100)}</sheetData>")


# -- what the blocks would hold, not only what the load needs -----------------------------------------


def _repeated_strings(width: int, n: int = 20, uses: int = 20) -> bytes:
    from tests.support.excel_workbooks import MAIN_NS, column_letters

    strings = "".join(f"<si><t>{i:03d}{'x' * width}</t></si>" for i in range(n))
    sst = f'<?xml version="1.0"?><sst xmlns="{MAIN_NS}">{strings}</sst>'
    rows = "".join(
        f'<row r="{i + 1}">' + "".join(f'<c r="{column_letters(c + 1)}{i + 1}" t="s"><v>{(i + c) % n}</v></c>' for c in range(uses)) + "</row>"
        for i in range(n)
    )
    return workbook_zip(f"<sheetData>{rows}</sheetData>", extra={"xl/sharedStrings.xml": sst})


def test_text_repeated_by_shared_strings_is_counted_per_use(monkeypatch: pytest.MonkeyPatch) -> None:
    data = _repeated_strings(10_000)  # 20 strings x 20 uses x 10 000 characters = 4 M bytes of block text, 8 M with the headers
    monkeypatch.setattr(module, "EXCEL_MAX_TEXT_CHARS", 6_000_000)
    with pytest.raises(ParseError) as raised:
        load(data)
    assert raised.value.message == user_errors.EXCEL_TOO_LARGE
    monkeypatch.setattr(module, "EXCEL_MAX_TEXT_CHARS", 12_000_000)
    assert load(data).workbook.text_bytes >= 4_000_000  # the same workbook loads under a larger allowance


def test_a_corrupt_workbook_keeps_the_loaders_own_error_and_a_non_zip_is_unreadable() -> None:
    from lxml import etree

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("[Content_Types].xml", "<broken")
    with pytest.raises(etree.XMLSyntaxError):  # as before the worker: not a stand-in RuntimeError
        load(out.getvalue())
    # Not a zip at all is the file's fault, not an outage: a 422, never a retried 500.
    with pytest.raises(ParseError) as raised:
        load(b"not a zip")
    assert raised.value.message == user_errors.UNREADABLE_FILE


# -- the running output budget: what the parent produces, charged as it is created ----------------------


def _loaded_with_budget(data: bytes, budget: int) -> ExcelParser:
    parser = load(data)  # the load-time estimate is not what these tests are about
    parser._output = module._OutputBudget(budget)
    return parser


def test_the_basic_blocks_stop_at_the_output_budget() -> None:
    parser = _loaded_with_budget(_small_xlsx_many_rows(300), 2_000)
    with pytest.raises(ParseError) as raised:
        parser._build_basic_block_container()
    assert raised.value.message == user_errors.EXCEL_TOO_LARGE


@pytest.mark.asyncio
async def test_the_indexing_blocks_stop_at_the_output_budget() -> None:
    parser = _loaded_with_budget(_small_xlsx_many_rows(300), 2_000)
    _stub_llm(parser)
    parser.get_table_summary = AsyncMock(return_value="summary")
    parser.get_rows_text = AsyncMock(side_effect=lambda rows, _s: ["x" * 500 for _ in rows])  # LLM output is charged
    with pytest.raises(ParseError) as raised:
        await parser.get_blocks_from_workbook(MagicMock())
    assert raised.value.message == user_errors.EXCEL_TOO_LARGE


def test_a_long_sheet_name_is_charged_for_every_block_it_is_written_into() -> None:
    # The name is stored once, but every block's data and the JSON of the container carry it.
    wb = openpyxl.Workbook()
    wb.active.title = "N" * 31  # Excel's limit; openpyxl's writer enforces it, the reader does not
    wb.active.append(["h"])
    for i in range(500):
        wb.active.append([i])
    out = io.BytesIO()
    wb.save(out)
    parser = _loaded_with_budget(out.getvalue(), 500 * 31 // 2)
    with pytest.raises(ParseError):
        parser._build_basic_block_container()


def test_the_output_budget_is_per_loaded_document() -> None:
    parser = load(_small_xlsx_many_rows(50))
    parser._build_basic_block_container()
    first_left = parser._output.left
    parser.load_workbook_from_binary(_small_xlsx_many_rows(50))
    assert parser._output.left > first_left  # a new load starts a new budget


def test_a_number_format_is_clamped_to_excels_limit_when_rendered() -> None:
    long_format = 'yyyy-mm-dd"' + "L" * 200_000 + '"'
    rendered = module.format_excel_datetime(dt.datetime(2024, 1, 2), long_format)
    assert len(rendered) < 1_000


def test_the_model_keeps_a_clamped_number_format() -> None:
    wb = openpyxl.Workbook()
    wb.active["A1"] = dt.datetime(2024, 1, 1)
    wb.active["A1"].number_format = 'yyyy-mm-dd"' + "L" * 5_000 + '"'
    out = io.BytesIO()
    wb.save(out)
    assert len(load(out.getvalue()).workbook["Sheet"].cell(1, 1).number_format) <= 255


def _small_xlsx_many_rows(n: int) -> bytes:
    wb = openpyxl.Workbook()
    wb.active.append(["name", "amount"])
    for i in range(n):
        wb.active.append([f"row {i}", i])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# -- structures count, not only text: a cell dict is ~1.1 KB however small the cell ----------------------


def _numbers_xlsx(rows: int, cols: int) -> bytes:
    wb = openpyxl.Workbook()
    for r in range(rows):
        wb.active.append([r + c for c in range(cols)])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_the_basic_parse_charges_every_cell_dict_not_only_text() -> None:
    cells = 200 * 50  # 10 000 one-digit numbers: no text to speak of
    data = _numbers_xlsx(200, 50)
    with pytest.raises(ParseError) as raised:
        _loaded_with_budget(data, cells * module._CELL_DICT_BYTES // 4)._build_basic_block_container()
    assert raised.value.message == user_errors.EXCEL_TOO_LARGE
    # the same workbook builds when the budget holds its cells, blocks and text
    ok = _loaded_with_budget(data, cells * module._CELL_DICT_BYTES * 2)
    assert ok._build_basic_block_container().blocks


@pytest.mark.asyncio
async def test_the_indexing_tables_charge_their_cell_dicts_and_give_them_back_per_sheet() -> None:
    wb = openpyxl.Workbook()
    for name in ("A", "B"):
        ws = wb.active if name == "A" else wb.create_sheet(name)
        ws.title = name
        ws.append([f"h{c}" for c in range(40)])
        for r in range(100):
            ws.append([r + c for c in range(40)])
    out = io.BytesIO()
    wb.save(out)
    one_sheet = 101 * 40 * module._CELL_DICT_BYTES  # a sheet's cell dicts, in bytes

    def parser(budget: int) -> ExcelParser:
        p = _loaded_with_budget(out.getvalue(), budget)
        _stub_llm(p)
        p.get_table_summary = AsyncMock(return_value="s")
        p.get_rows_text = AsyncMock(side_effect=lambda rows, _s: ["r" for _ in rows])
        return p

    with pytest.raises(ParseError):  # not even one sheet's cells fit
        await parser(one_sheet // 2).get_blocks_from_workbook(MagicMock())
    # one sheet's cells (plus rows and blocks) fit, two together would not: they are freed after each sheet
    container = await parser(int(one_sheet * 1.6)).get_blocks_from_workbook(MagicMock())
    assert len(container.blocks) == 200


def test_a_stray_cell_far_from_the_data_cannot_make_the_parent_build_millions_of_cells(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(module, "EXCEL_MAX_TEXT_CHARS", 8 * MB)  # a 24 MB budget: about 22 000 cell dicts
    wb = openpyxl.Workbook()
    wb.active["A1"] = "h"
    wb.active["A2"] = 1
    wb.active.cell(60_000, 300, 7)  # 18 M positions below the data
    out = io.BytesIO()
    wb.save(out)
    start = time.monotonic()
    with pytest.raises(ParseError) as raised:
        load(out.getvalue())._build_basic_block_container()
    assert raised.value.message == user_errors.EXCEL_TOO_LARGE
    assert time.monotonic() - start < 15  # refused after ~70 rows of 300 cells, not after 18 M cells
